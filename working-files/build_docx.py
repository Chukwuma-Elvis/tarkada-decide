"""
Builds .docx copies of the two submission PDFs directly from the same
markdown sources used by build_dossier.py / build_summary.py, since no
pandoc or LibreOffice is available on this machine to convert the PDFs
themselves. A small line-based markdown parser (headings, paragraphs,
**bold**, `code`, bullet/numbered lists, pipe tables) drives python-docx.
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from build_dossier import ORDER, BASE

SUBMISSION = os.path.join(BASE, "..", "submission")
RED = RGBColor(0x8B, 0x00, 0x00)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline_runs(paragraph, text):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def set_cell_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


def parse_table(doc, lines):
    rows = [l for l in lines if l.strip().startswith("|")]
    rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|$", r.strip())]
    cells_per_row = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    ncols = len(cells_per_row[0])
    table = doc.add_table(rows=len(cells_per_row), cols=ncols)
    table.style = "Table Grid"
    set_cell_borders(table)
    for ri, row in enumerate(cells_per_row):
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell_p = table.cell(ri, ci).paragraphs[0]
            add_inline_runs(cell_p, text)
            if ri == 0:
                for run in cell_p.runs:
                    run.bold = True
            for run in cell_p.runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def render_markdown(doc, text):
    lines = text.split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("| "):
            j = i
            while j < n and lines[j].strip().startswith("|"):
                j += 1
            parse_table(doc, lines[i:j])
            i = j
            continue
        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            add_inline_runs(h, stripped[4:])
            i += 1
            continue
        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            add_inline_runs(h, stripped[3:])
            i += 1
            continue
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            add_inline_runs(h, stripped[2:])
            i += 1
            continue
        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
            i += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\. ", "", stripped))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for level, size in [(1, 18), (2, 14), (3, 12)]:
        h = doc.styles[f"Heading {level}"]
        h.font.size = Pt(size)
        h.font.color.rgb = RED if level == 1 else RGBColor(0x33, 0x33, 0x33)
        h.font.bold = True
    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def build_dossier_docx():
    doc = Document()
    style_document(doc)

    title = doc.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Takarda Decision Dossier")
    run.font.size = Pt(28)
    run.font.color.rgb = RED
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "Architecture and decision record for the Takarda results-checking platform, "
        "prepared for the National Secondary Certificate Council contract."
    )
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    for _, label in ORDER:
        doc.add_paragraph(label)
    doc.add_page_break()

    for i, (relpath, label) in enumerate(ORDER):
        with open(os.path.join(BASE, relpath), encoding="utf-8") as f:
            text = f.read()
        render_markdown(doc, text)
        if i < len(ORDER) - 1:
            doc.add_page_break()

    out = os.path.join(SUBMISSION, "Takarda-Decision-Dossier.docx")
    doc.save(out)
    print("DOCX built:", out)


def build_summary_docx():
    doc = Document()
    style_document(doc)
    with open(os.path.join(BASE, "dossier", "00-one-page-summary.md"), encoding="utf-8") as f:
        text = f.read()
    render_markdown(doc, text)
    out = os.path.join(SUBMISSION, "One-Page-Summary.docx")
    doc.save(out)
    print("DOCX built:", out)


if __name__ == "__main__":
    build_dossier_docx()
    build_summary_docx()
